from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    BaseDocTemplate,
    Frame,
    KeepTogether,
    PageBreak,
    PageTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)


ROOT = Path(__file__).resolve().parents[1]
DOCX_OUT = ROOT / "output" / "docx" / "Nightingale_Technical_Brief.docx"
PDF_OUT = ROOT / "output" / "pdf" / "Nightingale_Technical_Brief.pdf"

BLUE = colors.HexColor("#153A5B")
TEAL = colors.HexColor("#168A84")
PALE = colors.HexColor("#EAF4F3")
INK = colors.HexColor("#18232E")
MUTED = colors.HexColor("#526170")
LINE = colors.HexColor("#CAD6DF")
RED = colors.HexColor("#9B2C2C")
AMBER = colors.HexColor("#936412")
GREEN = colors.HexColor("#216E4E")


SCENARIOS = [
    (1, "DOES NOT", "Phone-only identity", "Email/password blocks a phone/WhatsApp-only patient. Server sessions are sound, but OTP, consented delivery and assisted offline access do not exist."),
    (2, "PARTIAL", "Clinic isolation", "RBAC + scoped loaders + FKs/triggers/indexes cover tested routes. SQLite has no RLS; a missed route remains the first leak risk."),
    (3, "PARTIAL", "Logs and retention", "Application fields are allowlisted/scrubbed and edge access logs are off. Host, crash, third-party and Provider retention remain unestablished."),
    (4, "SURVIVES", "Redaction ordering", "Recursive redaction precedes the sole LLM egress; restore and exact anchoring are local and fail closed. A future direct Provider call is the regression risk."),
    (5, "PARTIAL", "Second clinic", "Bootstrap, settings and idempotent CSV import need no schema rewrite. Organisation verification, provisioning, per-clinic billing/support and RLS remain."),
    (6, "PARTIAL", "Trilingual consult", "UTF-8, manual speakers, attestations and exact spans work. Local ASR missed two Malay phrases; fallback is mainly English-keyword and clinical validity is unproven."),
    (7, "DOES NOT", "Minute-two allergy", "Voice is post-consult batch. Streaming chunks, incremental risk detection, latency targets and alert acknowledgement/escalation are a separate product path."),
    (8, "PARTIAL", "45-second hang", "A 30-second total deadline cancels the request and returns labelled fallback/unavailable. Cancellation was proven locally, not against a deliberately hung live Provider."),
    (9, "SURVIVES", "Provider 503", "Raw-first ingestion plus labelled deterministic fallback remains useful and may safely produce zero Highlights. It is not equivalent to multilingual Provider understanding."),
    (10, "SURVIVES", "Concurrent edits", "CAS yields one winner and a visible 409; snapshots, diff, revert-as-new-version and audit explain it. No distributed certification or live presence."),
    (11, "PARTIAL", "Delivery end-to-end", "Exact-version portal receipts distinguish view/acknowledge. There is no sender, Provider receipt, retry, bounce or escalation for email/SMS/WhatsApp."),
    (12, "PARTIAL", "Wrong patient copy", "Clinician publish/correct/withdraw gates visibility and preserves history. External copies cannot be recalled; medication accuracy still depends on human review."),
    (13, "SURVIVES", "Conflicting allergy", "Both exact sources remain and needs_review outranks routine items; the system does not pick truth. The parser is bounded English logic."),
    (14, "PARTIAL", "Importance meaning", "It is a retrieval-order score, not risk probability. Factors, source and Coverage Review make it falsifiable; clinical calibration/outcome evidence is absent."),
    (15, "PARTIAL", "Learning bias/fatigue", "Complete decisions, scope, caps, protections, replay/freeze/rollback exist. No real labels or fatigue policy; Shadow only and formal Glance stays base_only."),
    (16, "SURVIVES", "Edited source", "Version + exact Span + quote hash resolves historical text or fails closed. No dependency graph or production archive certification."),
    (17, "PARTIAL", "Cross-cutting synthesis", "The concluding checklist is integrated where bounded. Streaming, diarization, noisy-clinic validation, medical-reference confirmation, regeneration and production learning remain."),
]


def status_color(status: str):
    if status == "SURVIVES":
        return GREEN
    if status == "DOES NOT":
        return RED
    return AMBER


def add_page_number(canvas, doc):
    canvas.saveState()
    canvas.setStrokeColor(LINE)
    canvas.line(16 * mm, 13 * mm, 194 * mm, 13 * mm)
    canvas.setFont("Helvetica", 7)
    canvas.setFillColor(MUTED)
    canvas.drawString(16 * mm, 8.5 * mm, "Nightingale | Technical Brief | 03 Sep 2026")
    canvas.drawRightString(194 * mm, 8.5 * mm, f"{doc.page} / 3")
    canvas.restoreState()


def p(text, style):
    return Paragraph(text, style)


def build_pdf():
    styles = getSampleStyleSheet()
    title = ParagraphStyle("Title", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=24, leading=27, textColor=BLUE, spaceAfter=6)
    kicker = ParagraphStyle("Kicker", parent=styles["Normal"], fontName="Helvetica-Bold", fontSize=8, leading=10, textColor=TEAL, spaceAfter=8)
    h1 = ParagraphStyle("H1", parent=styles["Heading1"], fontName="Helvetica-Bold", fontSize=13, leading=15, textColor=BLUE, spaceBefore=5, spaceAfter=5)
    body = ParagraphStyle("Body", parent=styles["BodyText"], fontName="Helvetica", fontSize=8.3, leading=10.6, textColor=INK, spaceAfter=5)
    small = ParagraphStyle("Small", parent=body, fontSize=7.2, leading=8.6, spaceAfter=0)
    tiny = ParagraphStyle("Tiny", parent=body, fontSize=6.6, leading=7.9, spaceAfter=0)
    header = ParagraphStyle("Header", parent=tiny, fontName="Helvetica-Bold", textColor=colors.white)
    note = ParagraphStyle("Note", parent=body, fontSize=7.6, leading=9.4, textColor=MUTED, borderColor=LINE, borderWidth=0.6, borderPadding=6, backColor=colors.HexColor("#F7F9FA"))
    mono = ParagraphStyle("Mono", parent=body, fontName="Courier", fontSize=7.2, leading=9, textColor=INK, backColor=colors.HexColor("#F4F7F9"), borderPadding=7)

    doc = BaseDocTemplate(str(PDF_OUT), pagesize=A4, leftMargin=16 * mm, rightMargin=16 * mm, topMargin=14 * mm, bottomMargin=17 * mm, title="Nightingale Technical Brief", author="Yan Changhao")
    frame = Frame(doc.leftMargin, doc.bottomMargin, doc.width, doc.height, id="main")
    doc.addPageTemplates(PageTemplate(id="brief", frames=[frame], onPage=add_page_number))
    story = []

    story += [
        p("Nightingale", title),
        p("TECHNICAL BRIEF · FEEDBACK-INTEGRATED BUILD", kicker),
        p("<b>Status:</b> synthetic-data prototype — verified with explicit limits &nbsp;&nbsp; <b>Baseline:</b> main, 03 Sep 2026", body),
        p("The supplied clinic feedback numbers scenarios 1–16 and then adds a cross-cutting capability checklist. This brief treats that checklist as synthesis item 17; it does not invent a separately numbered source scenario.", note),
        p("1. Current build and integration model", h1),
        p("Nightingale is one longitudinal care record. A real-world <b>Event</b> is the Timeline unit; parallel <b>Artifacts</b> preserve raw, AI-scribed, clinician, staff and patient-facing representations without overwriting one another. Every derived Highlight binds to the source Artifact version, exact Span, verbatim quote and quote hash.", body),
    ]

    architecture = Table([
        [p("CLIENT PROJECTIONS", header), p("SERVER AUTHORITY", header), p("DURABLE / EXTERNAL", header)],
        [p("Timeline · what happened<br/>Glance · what matters now<br/>Patient View · what to do", small), p("RBAC + clinic scope<br/>state + human gates<br/>ranking + provenance", small), p("SQLCipher/SQLite<br/>metadata audit<br/>one redacted LLM egress", small)],
        [p("Patient -&gt; Event -&gt; Artifact -&gt; exact Span / version / hash", small), "", ""],
    ], colWidths=[54 * mm, 70 * mm, 54 * mm], rowHeights=[8 * mm, 22 * mm, 10 * mm])
    architecture.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), BLUE), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("BACKGROUND", (0, 1), (-1, 1), PALE), ("BACKGROUND", (0, 2), (-1, 2), colors.HexColor("#DCECEF")),
        ("SPAN", (0, 2), (-1, 2)), ("GRID", (0, 0), (-1, -1), 0.5, LINE),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 5),
    ]))
    story += [architecture, Spacer(1, 5), p("The server owns scope, authorship, state transitions, safety, Task authority, publication and provenance. AI may assist language only. Formal Glance stays deterministic <b>base_only</b>; learned ranking is Shadow-only.", body), p("2. Assumptions re-evaluated", h1)]

    assumptions = Table([
        [p("STILL STAND", small), p("NO LONGER STAND", small)],
        [p("• Event → Artifact → Span is the canonical model.<br/>• Raw source is immutable; human authority remains separate.<br/>• Server scope and patient-safe projection are mandatory.<br/>• Deterministic ranking precedes learned ranking.<br/>• The evidence is a synthetic, single-machine prototype.", small), p("• Email identity reaches every patient.<br/>• Redaction-before-model closes all privacy exits.<br/>• Batch ASR supports real-time alerts.<br/>• A generated link proves delivery.<br/>• Surfaced-only feedback is sufficient learning evidence.<br/>• Provider failure is always a returned error, not a hang.", small)],
    ], colWidths=[89 * mm, 89 * mm])
    assumptions.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (0, 0), colors.HexColor("#E8F3ED")), ("TEXTCOLOR", (0, 0), (0, 0), GREEN),
        ("BACKGROUND", (1, 0), (1, 0), colors.HexColor("#F7ECEB")), ("TEXTCOLOR", (1, 0), (1, 0), RED),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("GRID", (0, 0), (-1, -1), 0.5, LINE),
        ("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 6), ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    story += [assumptions, PageBreak(), p("Nightingale", kicker), p("3. Feedback traceability · scenarios 1–9", h1)]

    def scenario_table(rows):
        data = [[p("#", header), p("STATUS", header), p("FEEDBACK / BUILD INTEGRATION / FIRST FAILURE", header)]]
        for number, status, name, text in rows:
            data.append([p(str(number), tiny), p(status, ParagraphStyle(f"s{number}", parent=tiny, fontName="Helvetica-Bold", textColor=status_color(status))), p(f"<b>{name}.</b> {text}", tiny)])
        table = Table(data, colWidths=[8 * mm, 23 * mm, 147 * mm], repeatRows=1)
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), BLUE), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 0.35, LINE), ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("ALIGN", (0, 1), (0, -1), "CENTER"),
            ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ("TOPPADDING", (0, 0), (-1, -1), 3.5), ("BOTTOMPADDING", (0, 0), (-1, -1), 3.5),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFB")]),
        ]))
        return table

    story += [scenario_table(SCENARIOS[:9]), PageBreak(), p("Nightingale", kicker), p("4. Feedback traceability · scenarios 10–17", h1), scenario_table(SCENARIOS[9:]), Spacer(1, 5), p("Official numbered scenarios 1–16: <b>5 SURVIVE / 9 PARTIAL / 2 DOES NOT</b>. Cross-cutting synthesis item 17: <b>PARTIAL</b>.", note), p("5. What we tried, what failed, and where work stopped", h1)]

    failures = [
        ("DeepSeek protocol", "The first Anthropic-format run spent all 2,000 output tokens on reasoning and returned no final content. The OpenAI-compatible final-content path later produced one strict synthetic summary with 4/4 anchors; this is not clinical validation."),
        ("Local ASR", "Offline mechanics passed on supplied synthetic audio, but two Malay phrases were misrecognized and no diarization exists. We stopped at required human correction/attestation."),
        ("Tenant defense", "SQLite cannot supply database RLS. Scoped loaders, triggers, FKs, indexes and fault injection reduce blast radius; production tenancy awaits a database/platform decision."),
        ("Observed learning", "The compiler found zero eligible real-clinician pairs, so training correctly stopped. Frozen synthetic models remain Shadow-only and never serve Glance."),
        ("External boundaries", "A local hanging server proved timeout cancellation; no live Provider hang was induced. Messaging, host/crash retention, streaming alerts and recall of sent content were not built or simulated."),
    ]
    failure_table = Table([[p(a, small), p(b, small)] for a, b in failures], colWidths=[37 * mm, 141 * mm])
    failure_table.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.35, LINE), ("BACKGROUND", (0, 0), (0, -1), PALE),
        ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"), ("TEXTCOLOR", (0, 0), (0, -1), BLUE),
        ("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    story += [failure_table, Spacer(1, 5), p("6. Verification, evidence map, and foreseeable failures", h1)]

    metrics = Table([
        [p("BACKEND", header), p("FRONTEND", header), p("PROVENANCE / PROVIDER", header), p("PERFORMANCE", header)],
        [p("714 collected<br/><b>712 passed</b><br/>2 default ASR-input skips", small), p("Contract checks passed<br/><b>112-module build passed</b>", small), p("One synthetic DeepSeek run<br/><b>4/4 exact anchors</b>", small), p("Glance HTTP P95<br/><b>5.115 ms</b><br/>100 local warm samples", small)],
    ], colWidths=[44.5 * mm] * 4)
    metrics.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), BLUE), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("BACKGROUND", (0, 1), (-1, 1), PALE), ("GRID", (0, 0), (-1, -1), 0.5, LINE),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    story += [metrics, Spacer(1, 4), p("Scenario-linked automated evidence", h1), p("The README maps each numbered scenario to its primary backend tests; dated evidence files preserve exact commands, observations and non-claims.", body), p("Foreseeable first production failures", h1), p("Phone-only patient exclusion; ungoverned host/third-party retention; code-switching ASR error propagation; absence of in-consult alerts; undelivered external links; and over-trust in an uncalibrated retrieval score. Human gates, immutable raw sources, exact fail-closed provenance and deterministic serving limit—not eliminate—these risks.", body), p("Release boundary", h1), p("This is a <b>synthetic-data prototype</b>, not evidence of production medical safety, multilingual clinical accuracy, public-host security, production multi-tenancy, real-time alert delivery, a clinically valid learned ranker, regulatory compliance or external message delivery.", note)]

    doc.build(story)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def build_docx():
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.6)
    section.right_margin = Inches(0.6)
    styles = document.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(8.5)
    styles["Title"].font.name = "Arial"
    styles["Title"].font.size = Pt(24)
    styles["Title"].font.color.rgb = RGBColor(0x15, 0x3A, 0x5B)
    for name in ["Heading 1", "Heading 2"]:
        styles[name].font.name = "Arial"
        styles[name].font.color.rgb = RGBColor(0x15, 0x3A, 0x5B)
    styles["Heading 1"].font.size = Pt(13)

    title = document.add_paragraph(style="Title")
    title.add_run("Nightingale")
    sub = document.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = sub.add_run("TECHNICAL BRIEF · FEEDBACK-INTEGRATED BUILD")
    run.bold = True
    run.font.color.rgb = RGBColor(0x16, 0x8A, 0x84)
    document.add_paragraph("Status: synthetic-data prototype — verified with explicit limits | Baseline: main, 03 Sep 2026")
    document.add_paragraph("Scope note: the supplied feedback numbers scenarios 1–16 and then adds a cross-cutting checklist. This brief treats that checklist as synthesis item 17; it does not invent a separately numbered source scenario.")
    document.add_heading("1. Current build and integration model", level=1)
    document.add_paragraph("Nightingale is one longitudinal care record. A real-world Event is the Timeline unit; parallel Artifacts preserve raw, AI-scribed, clinician, staff and patient-facing representations. Every derived Highlight binds to the source Artifact version, exact Span, verbatim quote and quote hash.")
    document.add_paragraph("Patient → Event → Artifact → exact Span/version/hash\nBrowser/session → FastAPI → RBAC + scoped loaders → SQLCipher/SQLite\nRedacted content → one LLM egress → strict validation → exact source anchor")
    document.add_paragraph("The server owns scope, authorship, state, safety, Tasks, publication and provenance. AI assists language only. Formal Glance stays deterministic base_only; learned ranking is Shadow-only.")
    document.add_heading("2. Assumptions re-evaluated", level=1)
    document.add_paragraph("Still stand: Event → Artifact → Span; raw-source immutability; human authority; server scope; patient-safe projection; deterministic-before-learned ranking; synthetic single-machine boundary.")
    document.add_paragraph("No longer stand: email identity reaches every patient; redaction-before-model closes every privacy exit; batch ASR supports real-time alerts; a generated link proves delivery; surfaced-only feedback is sufficient; Provider failure is always a returned error.")
    document.add_page_break()
    document.add_heading("3. Feedback traceability · scenarios 1–9", level=1)
    add_scenario_docx_table(document, SCENARIOS[:9])
    document.add_page_break()
    document.add_heading("4. Feedback traceability · scenarios 10–17", level=1)
    add_scenario_docx_table(document, SCENARIOS[9:])
    document.add_paragraph("Official numbered scenarios 1–16: 5 SURVIVE / 9 PARTIAL / 2 DOES NOT. Cross-cutting synthesis item 17: PARTIAL.")
    document.add_heading("5. What we tried, what failed, and where work stopped", level=1)
    for name, detail in [
        ("DeepSeek protocol", "Anthropic-format exhausted 2,000 output tokens on reasoning and produced no final content. The final-content path later returned 4/4 exact anchors on one synthetic case; not clinical validation."),
        ("Local ASR", "Two Malay phrases were misrecognized and no diarization exists; the build stops at human correction/attestation."),
        ("Tenant defense", "SQLite has no RLS; defense-in-depth exists, while production tenancy awaits a database/platform decision."),
        ("Observed learning", "Zero eligible real-clinician pairs caused training to stop; synthetic models remain Shadow-only."),
        ("External boundaries", "Timeout cancellation was tested locally. Messaging, retention, streaming alerts and sent-copy recall were not built or simulated."),
    ]:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.add_run(f"{name}: ").bold = True
        paragraph.add_run(detail)
    document.add_heading("6. Verification, evidence, and foreseeable failures", level=1)
    document.add_paragraph("Current evidence: 714 backend tests collected, 712 passed, and two default skips requiring ignored local-ASR inputs; both passed when explicitly supplied. Frontend checks and the 112-module production build passed. Glance HTTP P95 was 5.115 ms over 100 local warm samples.")
    document.add_paragraph("Evidence map: identity/scope/log tests; redaction/timeout/fallback tests; consult/Voice/language tests; revision/publication/conflict tests; ranking/learning/provenance tests. The README provides the exact scenario-to-test index.")
    document.add_heading("Foreseeable first production failures", level=1)
    document.add_paragraph("Phone-only patient exclusion; host/third-party retention; code-switching ASR error propagation; absence of in-consult alerts; undelivered external links; and over-trust in an uncalibrated retrieval score.")
    document.add_heading("Release boundary", level=1)
    document.add_paragraph("This synthetic-data prototype does not establish production medical safety, multilingual clinical accuracy, public-host security, production multi-tenancy, real-time alert delivery, a clinically valid learned ranker, regulatory compliance, or external message delivery.")

    for section in document.sections:
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.add_run("Nightingale | Technical Brief | 03 Sep 2026")
    DOCX_OUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(DOCX_OUT)


def add_scenario_docx_table(document, rows):
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    headers = table.rows[0].cells
    headers[0].text = "#"
    headers[1].text = "Status"
    headers[2].text = "Feedback / build integration / first failure"
    for cell in headers:
        set_cell_shading(cell, "153A5B")
        for run in cell.paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.bold = True
    for number, status, name, detail in rows:
        cells = table.add_row().cells
        cells[0].text = str(number)
        cells[1].text = status
        cells[2].text = f"{name}. {detail}"
        for run in cells[1].paragraphs[0].runs:
            run.bold = True
        for cell in cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(7)
    table.columns[0].width = Inches(0.35)
    table.columns[1].width = Inches(0.9)
    table.columns[2].width = Inches(5.9)


if __name__ == "__main__":
    PDF_OUT.parent.mkdir(parents=True, exist_ok=True)
    build_pdf()
    build_docx()
    print(PDF_OUT)
    print(DOCX_OUT)
